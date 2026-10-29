"""Lecture d'entree et ecriture sans reconstruire les feuilles template."""
from __future__ import annotations

import re
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

SHEETS = {"issuer": "Template Import Titres", "fund": "Template Import Fonds"}


def valid_isin(value: str) -> bool:
    """Controle syntaxique et cle de controle Luhn ISIN."""
    if not re.fullmatch(r"[A-Z]{2}[A-Z0-9]{9}[0-9]", value):
        return False
    digits = "".join(str(ord(c) - 55) if c.isalpha() else c for c in value)
    total = 0
    for index, char in enumerate(reversed(digits)):
        number = int(char) * (2 if index % 2 else 1)
        total += number // 10 + number % 10
    return total % 10 == 0


def headers(template: Path) -> dict[str, list[str]]:
    workbook = load_workbook(template)
    try:
        result = {entity: [cell.value for cell in workbook[name][1]]
                  for entity, name in SHEETS.items()}
        for entity, names in result.items():
            if any(not isinstance(name, str) or not name for name in names):
                raise ValueError(f"En-tete vide ou non textuel dans {entity}.")
            if len(names) != len(set(names)):
                raise ValueError(f"En-tetes dupliques dans le template {entity}.")
        return result
    finally:
        workbook.close()


def read_input(path: Path, sheets: dict, isin_column: str) -> dict[str, list[dict]]:
    workbook = load_workbook(path, data_only=False)
    result = {}
    try:
        for entity, name in sheets.items():
            result[entity] = []
            if name not in workbook.sheetnames:
                continue
            values = iter(workbook[name].values)
            columns = next(values, ())
            if isin_column not in columns:
                raise ValueError(f"Colonne {isin_column} absente de {name}.")
            if len([c for c in columns if c]) != len(set(c for c in columns if c)):
                raise ValueError(f"En-tetes dupliques dans {name}.")
            for row_number, values_row in enumerate(values, start=2):
                if not any(v is not None for v in values_row):
                    continue
                if any(isinstance(v, str) and v.startswith("=") for v in values_row):
                    raise ValueError(f"Formule interdite dans {name}, ligne {row_number}: fournir des valeurs.")
                row = dict(zip(columns, values_row))
                row["_row"] = row_number
                row["_isin"] = str(row.get(isin_column) or "").strip().upper()
                result[entity].append(row)
        if not any(result.values()):
            raise ValueError("Aucune ligne d'entree dans les feuilles configurees.")
        return result
    finally:
        workbook.close()


def export(template: Path, output: Path, rows: dict[str, list[dict]]) -> None:
    """Remplit une copie du template; conserve en-tetes, styles et dimensions."""
    if template.resolve() == output.resolve():
        raise ValueError("Ne pas ecraser le template.")
    workbook = load_workbook(template)
    temporary = output.with_name(output.stem + ".tmp.xlsx")
    try:
        for name in list(workbook.sheetnames):
            if name not in SHEETS.values():
                del workbook[name]
        for entity, name in SHEETS.items():
            sheet = workbook[name]
            columns = [cell.value for cell in sheet[1]]
            if any(cell.value is not None for row in sheet.iter_rows(min_row=2) for cell in row):
                raise ValueError(f"Le template {name} doit etre vide sous les en-tetes.")
            for index, row in enumerate(rows.get(entity, []), start=2):
                if sheet.row_dimensions[2].height is not None:
                    sheet.row_dimensions[index].height = sheet.row_dimensions[2].height
                for column, header in enumerate(columns, start=1):
                    cell = sheet.cell(index, column)
                    if sheet.cell(2, column).has_style:
                        cell._style = copy(sheet.cell(2, column)._style)
                    value = row.get(header)
                    cell.value = value
                    if isinstance(value, str):
                        # Texte force: pas d'execution de formule importee.
                        cell.data_type = "s"
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(temporary)
        temporary.replace(output)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)

"""Fixtures artificielles uniquement pour tests, jamais exportees comme donnees MSCI."""
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from msci_enrichment.client import MSCIClient
from msci_enrichment.excel import SHEETS, export, headers, valid_isin
from msci_enrichment.pipeline import flatten, run

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates/mapping.xlsx"


def test_isin_validation():
    assert valid_isin("US0378331005")
    assert not valid_isin("US0378331006")
    assert not valid_isin("")


def test_mapping_matches_template():
    mapping = json.loads((ROOT / "config/mapping.json").read_text(encoding="utf-8"))
    for entity, columns in headers(TEMPLATE).items():
        assert list(mapping[entity]) == columns
    assert mapping["issuer"]["ITR"]["source"] == "ITR"
    assert mapping["fund"]["FUND_IMPLIED_TEMP_RISE"]["source"] == "FUND_IMPLIED_TEMP_RISE"


def test_export_preserves_headers_styles_and_dimensions(tmp_path):
    output = tmp_path / "output.xlsx"
    export(TEMPLATE, output, {"issuer": [{"ISSUER_NAME": "=not_a_formula", "ESG_SCORE": 0}]})
    original, result = load_workbook(TEMPLATE), load_workbook(output)
    assert result.sheetnames == list(SHEETS.values())
    for name in SHEETS.values():
        assert [c.value for c in original[name][1]] == [c.value for c in result[name][1]]
        assert original[name].freeze_panes == result[name].freeze_panes
        for before, after in zip(original[name][1], result[name][1]):
            assert before._style == after._style
        for key, dimension in original[name].column_dimensions.items():
            assert dimension.width == result[name].column_dimensions[key].width
    assert result[SHEETS["issuer"]]["A2"].data_type == "s"
    assert result[SHEETS["issuer"]]["L2"].value == 0


def test_flatten_ambiguity():
    with pytest.raises(ValueError):
        flatten({"issuer_metadata": [{"ISSUERID": "a"}, {"ISSUERID": "b"}]})


def test_batch_retrieval():
    client = MSCIClient("test", "test")
    calls = []
    def request(method, endpoint, **kwargs):
        calls.append(kwargs["json"])
        return {"result": {"response_metadata": {
            "total_number_of_batches": 3, "data_request_id": "test-request"}, "data": []}}
    client.request = request
    assert len(client.batches("/issuers", {"issuer_identifier_list": ["test"]})) == 3
    assert calls[2] == {"data_request_id": "test-request", "batch_id": 3}
    client.close()


class FakeClient:
    def factors(self, entity):
        return {"IVA_COMPANY_RATING": {"factor_type": "issuer"}}

    def batches(self, endpoint, payload):
        return [{"result": {"response_metadata": {}, "data": [{
            "requested_id": "US0378331005",
            "issuer_metadata": [{"ISSUER_NAME": "TEST FIXTURE", "ISSUERID": "TEST-ID"}],
            "factors": [{"name": "IVA_COMPANY_RATING", "value": "AA"}],
        }]}}]


def test_pipeline_preserves_duplicates_invalid_and_blanks(tmp_path):
    input_file = tmp_path / "input.xlsx"
    w = Workbook()
    s = w.active
    s.title = "Titres"
    s.append(["ISIN", "Client_ID"])
    s.append(["US0378331005", "first"])
    s.append(["US0378331005", "second"])
    s.append(["invalid", "third"])
    w.save(input_file)
    settings = json.loads((ROOT / "config/settings.json").read_text())
    settings.update(template=str(TEMPLATE), mapping=str(ROOT / "config/mapping.json"),
                    input=str(input_file), output=str(tmp_path / "output.xlsx"),
                    diagnostics=str(tmp_path / "diagnostics.json"))
    report = run(FakeClient(), settings, tmp_path)
    w = load_workbook(tmp_path / "output.xlsx")
    s = w[SHEETS["issuer"]]
    assert s.max_row == 4
    assert s["B2"].value == "first" and s["B3"].value == "second"
    assert s["K2"].value == "AA" and s["K3"].value == "AA"
    assert s["K4"].value is None
    assert report["rows"][-1]["state"] == "invalid_isin"
    assert s["AE2"].value is None  # ITR remains blank.


def test_fund_resolution_and_zero(tmp_path):
    class FundClient:
        def factors(self, entity):
            return {"FUND_IMPLIED_TEMP_RISE": {"factor_type": "fund"}}

        def batches(self, endpoint, payload):
            if endpoint == "/coverage":
                data = [{"ISIN": "US0378331005", "MDSUID": "UFD-TEST", "FUND_NAME": "FIXTURE"}]
            else:
                assert payload["fund_identifier_list"] == ["UFD-TEST"]
                data = [{"requested_id": "UFD-TEST", "factors": [
                    {"name": "FUND_IMPLIED_TEMP_RISE", "value": 0}]}]
            return [{"result": {"data": data, "response_metadata": {}}}]

    w = Workbook()
    w.active.title = "Fonds"
    w.active.append(["ISIN"])
    w.active.append(["US0378331005"])
    w.save(tmp_path / "input.xlsx")
    settings = json.loads((ROOT / "config/settings.json").read_text())
    settings.update(template=str(TEMPLATE), mapping=str(ROOT / "config/mapping.json"),
                    input=str(tmp_path / "input.xlsx"), output=str(tmp_path / "output.xlsx"),
                    diagnostics=str(tmp_path / "diagnostics.json"), fund_products=["TEST-PRODUCT"])
    run(FundClient(), settings, tmp_path)
    sheet = load_workbook(tmp_path / "output.xlsx")[SHEETS["fund"]]
    assert sheet["B2"].value == "FIXTURE"
    column = headers(TEMPLATE)["fund"].index("FUND_IMPLIED_TEMP_RISE") + 1
    assert sheet.cell(2, column).value == 0


def test_token_reuse_and_401_refresh():
    class Response:
        def __init__(self, status, body):
            self.status_code, self.body, self.headers = status, body, {}

        def json(self):
            return self.body

    class Session:
        def __init__(self):
            self.auth_count, self.calls = 0, 0

        def post(self, *args, **kwargs):
            self.auth_count += 1
            return Response(200, {"access_token": "TEST", "expires_in": 86400})

        def request(self, *args, **kwargs):
            self.calls += 1
            return Response(401, {}) if self.calls == 2 else Response(200, {"result": {}})

        def close(self):
            pass

    session = Session()
    client = MSCIClient("TEST", "TEST", session=session)
    client.request("GET", "/metadata/factors")
    client.request("GET", "/metadata/factors")
    client.request("GET", "/metadata/factors")
    assert session.auth_count == 2
    assert session.calls == 4

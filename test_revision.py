"""Regressions issues des tests utilisateur et controles de transport.

Les reponses minimales reprennent les structures observees; aucune connexion reelle.
"""
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from msci_enrichment.client import MSCIClient, MSCIError
from msci_enrichment.diagnostics import describe_response, row_state
from msci_enrichment.pipeline import flatten, run

ROOT = Path(__file__).resolve().parents[1]


def settings_for(tmp_path):
    w = Workbook()
    w.active.title = 'Titres'
    w.active.append(['ISIN'])
    w.active.append(['GB00BTK05J60'])
    w.save(tmp_path / 'input.xlsx')
    settings = json.loads((ROOT / 'config/settings.json').read_text())
    settings.update(template=str(ROOT / 'templates/mapping.xlsx'),
                    mapping=str(ROOT / 'config/mapping.json'),
                    input=str(tmp_path / 'input.xlsx'), output=str(tmp_path / 'output.xlsx'),
                    diagnostics=str(tmp_path / 'diagnostics.json'))
    return settings


class ReferenceOnlyClient:
    def factors(self, entity):
        return {'IVA_COMPANY_RATING': {'factor_type': 'issuer'}}

    def batches(self, endpoint, payload):
        # Regression principale: les references doivent etre demandees explicitement.
        assert {'ISSUER_NAME', 'ISSUER_TICKER', 'IVA_COMPANY_RATING'} <= set(payload['factor_name_list'])
        return [{'status': 'OK', 'code': 200, 'trace_id': 'test-trace',
                 'messages': ['IVA_COMPANY_RATING not recognized or not available to this user'],
                 'result': {'response_metadata': {
                     'resolved_factors': ['ISSUER_NAME', 'ISSUER_TICKER'],
                     'total_number_of_unresolved_identifiers': 0},
                     'data': [{'requested_id': 'GB00BTK05J60', 'issuer_metadata': [{
                         'ISSUERID': 'IID000000002123931', 'ISSUER_NAME': 'ANGLO AMERICAN PLC',
                         'ISSUER_TICKER': 'AAL', 'ISSUER_ISIN': 'GB00BTK05J60'}],
                         'factors': [{'name': 'ISSUER_NAME', 'value': 'ANGLO AMERICAN PLC'},
                                     {'name': 'ISSUER_TICKER', 'value': 'AAL'}]}]}}]


def test_reference_only_is_not_esg_success(tmp_path):
    report = run(ReferenceOnlyClient(), settings_for(tmp_path), tmp_path)
    assert report['rows'][0]['state'] == 'reference_only'
    assert report['summary']['rows_with_esg_values'] == 0
    assert report['status'] == 'completed_with_warnings'
    assert report['api'][0]['requested_but_not_resolved'] == ['IVA_COMPANY_RATING']
    assert report['api'][0]['trace_id'] == 'test-trace'
    s = load_workbook(tmp_path / 'output.xlsx')['Template Import Titres']
    assert s['A2'].value == 'ANGLO AMERICAN PLC'
    assert s['D2'].value == 'AAL'
    assert s['K2'].value is None


def test_failure_keeps_diagnostic(tmp_path):
    class FailedClient:
        def factors(self, entity):
            raise MSCIError('HTTP 403')
    with pytest.raises(MSCIError):
        run(FailedClient(), settings_for(tmp_path), tmp_path)
    report = json.loads((tmp_path / 'diagnostics.json').read_text())
    assert report['status'] == 'failed'
    assert not (tmp_path / 'output.xlsx').exists()


@pytest.mark.parametrize('value', [0, False, '', None])
def test_scalar_values_are_not_imputed(value):
    assert flatten({'factors': [{'name': 'TEST', 'value': value}]})['TEST'] == value


def test_conflicting_metadata_is_not_silently_overwritten():
    with pytest.raises(ValueError):
        flatten({'issuer_metadata': [{'ISSUER_NAME': 'A'}],
                 'factors': [{'name': 'ISSUER_NAME', 'value': 'B'}]})


def test_timeseries_is_not_arbitrarily_reduced():
    with pytest.raises(ValueError):
        flatten({'factors': [{'name': 'TEST', 'data_values': [{'value': 1}]}]})


def test_empty_api_response_diagnostic():
    response = {'result': {'response_metadata': {'resolved_factors': ['ISSUER_CNTRY_DOMICILE'],
                'identifiers_data_not_available': ['GB00BTK05J60']}, 'data': []},
                'messages': ['IVA_COMPANY_RATING ignored']}
    info = describe_response(response, ['IVA_COMPANY_RATING', 'ISSUER_CNTRY_DOMICILE'])
    assert info['rows_returned'] == 0
    assert info['requested_but_not_resolved'] == ['IVA_COMPANY_RATING']
    assert row_state(True, False, False, False, 0, []) == 'no_data_returned'


class Response:
    def __init__(self, status, body=None, headers=None):
        self.status_code = status
        self.body = body or {'result': {}}
        self.headers = headers or {}

    def json(self):
        return self.body


class Session:
    def __init__(self, responses):
        self.responses = iter(responses)
        self.auth_count = 0

    def post(self, *args, **kwargs):
        self.auth_count += 1
        return Response(200, {'access_token': 'SECRET_TEST_TOKEN', 'expires_in': 86400})

    def request(self, *args, **kwargs):
        assert kwargs['allow_redirects'] is False
        return next(self.responses)

    def close(self):
        pass


def test_403_is_not_retried_and_no_token_in_audit():
    client = MSCIClient('ID', 'SECRET', session=Session([Response(403)]))
    with pytest.raises(MSCIError):
        client.request('POST', '/issuers', json={})
    assert 'SECRET' not in json.dumps(client.audit)
    assert client.audit[0]['http_status'] == 403


def test_429_retry_after(monkeypatch):
    delays = []
    monkeypatch.setattr('msci_enrichment.client.time.sleep', delays.append)
    client = MSCIClient('ID', 'SECRET', session=Session([
        Response(429, headers={'Retry-After': '2'}), Response(200)]))
    client.request('GET', '/metadata/factors')
    assert delays == [2]
    assert client.session.auth_count == 1


def test_coverage_503_load_retry(monkeypatch):
    delays = []
    monkeypatch.setattr('msci_enrichment.client.time.sleep', delays.append)
    client = MSCIClient('ID', 'SECRET', session=Session([Response(503), Response(200)]))
    client.request('POST', '/coverage', json={})
    assert delays == [300]


def test_http_200_with_application_error_is_failure():
    client = MSCIClient('ID', 'SECRET', session=Session([
        Response(200, {'code': 403, 'result': {}})]))
    with pytest.raises(MSCIError):
        client.request('POST', '/issuers', json={})


def test_coverage_batch_keeps_entity_type():
    client = MSCIClient('ID', 'SECRET')
    calls = []
    def request(method, endpoint, **kwargs):
        calls.append(kwargs['json'])
        return {'result': {'response_metadata': {'total_number_of_batches': 2,
                          'data_request_id': 'TEST', 'current_batch': len(calls)}, 'data': []}}
    client.request = request
    client.batches('/coverage', {'entity_type': 'fund', 'product': 'TEST'})
    assert calls[1] == {'entity_type': 'fund', 'data_request_id': 'TEST', 'batch_id': 2}
    client.close()


def test_paths_cannot_overwrite_input(tmp_path):
    settings = settings_for(tmp_path)
    settings['output'] = settings['input']
    with pytest.raises(ValueError):
        run(ReferenceOnlyClient(), settings, tmp_path)


def test_input_formulas_are_rejected(tmp_path):
    from msci_enrichment.excel import read_input
    w = Workbook()
    w.active.title = 'Titres'
    w.active.append(['ISIN'])
    w.active.append(['="GB00BTK05J60"'])
    w.save(tmp_path / 'input.xlsx')
    with pytest.raises(ValueError, match='Formule interdite'):
        read_input(tmp_path / 'input.xlsx', {'issuer': 'Titres'}, 'ISIN')


def test_unknown_endpoint_cannot_receive_token():
    client = MSCIClient('ID', 'SECRET')
    with pytest.raises(ValueError):
        client.request('GET', 'https://other.invalid')
    client.close()
